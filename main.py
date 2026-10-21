import pymupdf
import pytesseract
from pytesseract import Output
import pandas as pd
import re
import cv2
import numpy as np
import os
import json
import hashlib
from difflib import SequenceMatcher
from openpyxl import load_workbook
pytesseract.pytesseract.tesseract_cmd = r'C:\Program Files\Tesseract-OCR\tesseract.exe'

PASTA_PDFS = "PDFs"
ARQUIVO_SAIDA_XLSX = "saida_hibrida.xlsx"
ARQUIVO_SAIDA_PDF = "pdfs_unificados.pdf"
PASTA_CACHE = "cache_header"
DPI = 300
CROP_HEADER_RATIO = 0.5

PRODUTOS_ESPERADOS = [
    "ABOBORA", "BANANA", "BATATA DOCE", "BATATA INGLESA", "CARA",
    "CEBOLA", "CENOURA", "CHUCHU", "COENTRO", "MACAXEIRA",
    "MAMAO", "MELANCIA", "MELAO", "OVOS", "PIMENTAO", "TOMATE",
]
PRODUTOS_SEM_KG = {"OVOS"}
VARIACOES_KG = {"KG", "K6", "K.G", "K.G.", "K G", "KS"}
TOLERANCIA_LINHA_PDF = 3.0
TOLERANCIA_LINHA_OCR = 15
HEADER_KEYWORDS = {
    "QTD", "QTDE", "DESCRICAO", "DESCRIÇÃO", "CODIGO", "CÓDIGO",
    "U.M.", "U.M", "UM", "PESO", "PRODUTO",
}

# Larguras máximas por coluna na tabela do console
LARGURA_MAX_COLUNA = 45


# ==========================================
# UTILITÁRIOS GERAIS
# ==========================================
def ordenar_numericamente(lista):
    def chave(a):
        nome = os.path.splitext(a)[0]
        try:
            return (0, int(nome))
        except ValueError:
            return (1, nome)
    return sorted(lista, key=chave)


def imprimir_tabela(df, largura_max=LARGURA_MAX_COLUNA):
    """Imprime o DataFrame como tabela com larguras ajustadas ao conteúdo."""
    if df.empty:
        print("(vazio)")
        return

    colunas = list(df.columns)
    larguras = {}
    for col in colunas:
        larguras[col] = len(str(col))
        for val in df[col].astype(str):
            tamanho = min(len(val), largura_max)
            if tamanho > larguras[col]:
                larguras[col] = tamanho

    # Cabeçalho
    header = " | ".join(str(c).ljust(larguras[c]) for c in colunas)
    print(header)
    print("-" * len(header))

    # Linhas
    for _, row in df.iterrows():
        celulas = []
        for c in colunas:
            val = str(row[c])
            if len(val) > largura_max:
                val = val[: largura_max - 1] + "…"
            celulas.append(val.ljust(larguras[c]))
        print(" | ".join(celulas))


def autosize_excel(caminho):
    """Ajusta a largura das colunas do Excel ao conteúdo."""
    try:
        wb = load_workbook(caminho)
        ws = wb.active
        for col_cells in ws.iter_cols():
            max_len = 0
            col_letter = col_cells[0].column_letter
            for cell in col_cells:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = min(max_len + 2, 60)
        wb.save(caminho)
    except Exception as e:
        print(f"⚠️ Não foi possível ajustar largura do Excel: {e}")


# ==========================================
# CAMADA 1: PyMuPDF (produtos)
# ==========================================
def extrair_palavras_pdf(pagina):
    return [
        {"texto": w[4], "x0": w[0], "y0": w[1], "x1": w[2], "y1": w[3]}
        for w in pagina.get_text("words")
    ]


def agrupar_linhas_pdf(palavras, tol=TOLERANCIA_LINHA_PDF):
    linhas = {}
    for p in palavras:
        y = p["y0"]
        encaixou = False
        for yk in linhas:
            if abs(yk - y) < tol:
                linhas[yk].append(p)
                encaixou = True
                break
        if not encaixou:
            linhas[y] = [p]
    return dict(sorted(linhas.items()))


def encontrar_produto_fuzzy(texto, produtos, threshold=0.75, min_word_score=0.70):
    if not texto:
        return None
    limpo = re.sub(r'[^A-ZÀ-Ý\s]', ' ', texto.upper())
    tokens = [t for t in limpo.split() if t]
    if not tokens:
        return None
    junto = " ".join(tokens)
    melhor, melhor_score = None, 0
    for prod in produtos:
        if prod in junto:
            return prod
        scores = []
        for pp in prod.split():
            melhor_local = max(
                (SequenceMatcher(None, pp, t).ratio() for t in tokens),
                default=0,
            )
            scores.append(melhor_local)
        if min(scores) < min_word_score:
            continue
        medio = sum(scores) / len(scores)
        if medio > melhor_score:
            melhor_score, melhor = medio, prod
    return melhor if melhor_score >= threshold else None


def separar_colados(produto, produtos):
    nome = produto.upper().replace(" KG", "").strip()
    for prod in produtos:
        if nome.startswith(prod) and nome != prod:
            resto = nome[len(prod):].strip()
            if resto in produtos:
                return [f"{prod} KG", f"{resto} KG"]
    return [produto]


def encontrar_cabecalho_pdf(linhas):
    melhor_y, melhor_score = None, 0
    for y, palavras in linhas.items():
        textos = [p["texto"].upper().strip('.,;:') for p in palavras]
        score = sum(1 for t in textos if any(k in t for k in HEADER_KEYWORDS))
        if score >= 2 and score > melhor_score:
            melhor_score, melhor_y = score, y
    return melhor_y


def extrair_produtos_pdf(pagina):
    palavras = extrair_palavras_pdf(pagina)
    if not palavras:
        return [], "sem-texto"

    linhas = agrupar_linhas_pdf(palavras)
    texto_pagina = " ".join(p["texto"] for p in palavras)

    if "CHECKLIST" in texto_pagina.upper():
        return [], "checklist"

    y_cabecalho = encontrar_cabecalho_pdf(linhas)
    y_fim = float('inf')
    if y_cabecalho is not None:
        for y, palavras_linha in linhas.items():
            if y > y_cabecalho:
                txt = " ".join(p["texto"] for p in palavras_linha)
                if "Totais" in txt or "Observação" in txt:
                    y_fim = y
                    break

    modo = "cabecalho" if y_cabecalho is not None else "varredura"
    produtos = []

    for y, palavras_linha in linhas.items():
        if y_cabecalho is not None:
            if y <= y_cabecalho or y > y_fim:
                continue

        ordenadas = sorted(palavras_linha, key=lambda p: p["x0"])
        texto_linha = " ".join(p["texto"] for p in ordenadas)

        if any(h in texto_linha for h in ["Descrição", "U.M.", "Código", "CÓDIGO", "QTD"]):
            continue

        idx_kg = -1
        for i, p in enumerate(ordenadas):
            if p["texto"].upper().strip('.,;:') in VARIACOES_KG:
                idx_kg = i
                break

        if idx_kg != -1:
            candidato = " ".join(p["texto"] for p in ordenadas[:idx_kg + 1])
            prod = encontrar_produto_fuzzy(candidato, PRODUTOS_ESPERADOS)
            if prod:
                for pf in separar_colados(f"{prod} KG", PRODUTOS_ESPERADOS):
                    if pf not in produtos:
                        produtos.append(pf)
            continue

        for sem_kg in PRODUTOS_SEM_KG:
            if re.search(r'\b' + re.escape(sem_kg) + r'\b', texto_linha.upper()):
                if sem_kg not in produtos:
                    produtos.append(sem_kg)
                break

    return produtos, modo


# ==========================================
# CAMADA 2: Tesseract (data + lugar)
# ==========================================
def agrupar_linhas_ocr(dados_ocr, tol=TOLERANCIA_LINHA_OCR):
    linhas = {}
    for i in range(len(dados_ocr['text'])):
        texto = dados_ocr['text'][i].strip()
        if not texto:
            continue
        y = dados_ocr['top'][i]
        encaixou = False
        for yk in linhas:
            if abs(yk - y) < tol:
                linhas[yk].append({"texto": texto, "left": dados_ocr['left'][i]})
                encaixou = True
                break
        if not encaixou:
            linhas[y] = [{"texto": texto, "left": dados_ocr['left'][i]}]
    return dict(sorted(linhas.items()))


def extrair_data_ocr(texto):
    m = re.search(r'\bData\s*:?\s*(\d{1,2})[/\-\.](\d{1,2})[/\-\.](\d{4})', texto, re.IGNORECASE)
    if m:
        d, mo, a = m.group(1), m.group(2), m.group(3)
        return f"{int(d):02d}/{int(mo):02d}/{a}", f"{int(mo):02d}"
    m = re.search(r'(\d{1,2})/(\d{1,2})/(\d{4})', texto)
    if m:
        d, mo, a = m.group(1), m.group(2), m.group(3)
        return f"{int(d):02d}/{int(mo):02d}/{a}", f"{int(mo):02d}"
    return None, None


def extrair_lugar_ocr(dados_ocr):
    linhas = agrupar_linhas_ocr(dados_ocr)
    for y, palavras in linhas.items():
        ordenadas = sorted(palavras, key=lambda p: p['left'])
        textos = [p['texto'] for p in ordenadas]
        if "Nome:" in textos:
            idx = textos.index("Nome:")
            partes = []
            for t in textos[idx + 1:]:
                if any(lbl in t for lbl in [
                    "L.F:", "Fone:", "Cnpj:", "Cep:", "U.F:", "L.F",
                    "Fone", "L.F.", "IE:", "CNPJ:", "C.N.P.J", "Endereço:"
                ]):
                    break
                partes.append(t)
            nome = " ".join(partes).strip()
            return nome if nome else "NÃO IDENTIFICADO"
    return "NÃO IDENTIFICADO"


def extrair_header_ocr(pagina, pdf_hash, num_pagina):
    os.makedirs(PASTA_CACHE, exist_ok=True)
    cache_file = os.path.join(PASTA_CACHE, f"{pdf_hash}_p{num_pagina:04d}_header.json")

    if os.path.exists(cache_file):
        with open(cache_file, 'r', encoding='utf-8') as f:
            dados_ocr = json.load(f)
    else:
        rect = pagina.rect
        clip = pymupdf.Rect(0, 0, rect.width, rect.height * CROP_HEADER_RATIO)
        pix = pagina.get_pixmap(clip=clip, dpi=DPI)
        img = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.h, pix.w, pix.n)
        img_cinza = cv2.cvtColor(img, cv2.COLOR_RGB2GRAY)
        dados_ocr = pytesseract.image_to_data(
            img_cinza, lang='por', output_type=Output.DICT, config='--psm 6'
        )
        with open(cache_file, 'w', encoding='utf-8') as f:
            json.dump(dados_ocr, f, ensure_ascii=False)

    texto = " ".join(t for t in dados_ocr['text'] if t.strip())
    data, mes = extrair_data_ocr(texto)
    lugar = extrair_lugar_ocr(dados_ocr)
    return data, mes, lugar


# ==========================================
# UNIR PDFs
# ==========================================
def unir_pdfs(caminhos, saida):
    doc = pymupdf.open()
    for c in caminhos:
        t = pymupdf.open(c)
        doc.insert_pdf(t)
        t.close()
    doc.save(saida)
    doc.close()
    print(f"\n📎 PDF unificado: {saida}")


# ==========================================
# EXECUÇÃO PRINCIPAL
# ==========================================
def main():
    if not os.path.isdir(PASTA_PDFS):
        print(f"❌ Pasta '{PASTA_PDFS}' não encontrada.")
        return

    arquivos = [f for f in os.listdir(PASTA_PDFS) if f.lower().endswith('.pdf')]
    if not arquivos:
        print(f"❌ Nenhum PDF em '{PASTA_PDFS}'.")
        return

    arquivos = ordenar_numericamente(arquivos)
    print(f"📂 {len(arquivos)} PDFs encontrados.\n")

    todos_dados = []
    caminhos = []

    for idx, nome in enumerate(arquivos, start=1):
        caminho = os.path.join(PASTA_PDFS, nome)
        caminhos.append(caminho)
        print(f"[{idx}/{len(arquivos)}] {nome}")

        try:
            doc = pymupdf.open(caminho)
            pdf_hash = hashlib.md5(open(caminho, 'rb').read()).hexdigest()[:10]

            for num_pag, pagina in enumerate(doc, start=1):
                produtos, modo = extrair_produtos_pdf(pagina)
                data, mes, lugar = extrair_header_ocr(pagina, pdf_hash, num_pag)

                status = data if data else "❌ SEM DATA"
                print(f"    pág {num_pag}: {len(produtos):2d} produtos (modo={modo}), data={status}")

                for prod in produtos:
                    todos_dados.append({
                        "arquivo": nome,
                        "lugar": lugar,
                        "produto": prod,
                        "data": data,
                        "mes": mes,
                    })

            doc.close()
        except Exception as e:
            print(f"  ❌ Erro: {e}")

    if todos_dados:
        df = pd.DataFrame(todos_dados)
        df = df[["arquivo", "lugar", "produto", "data", "mes"]]
        df.to_excel(ARQUIVO_SAIDA_XLSX, index=False)
        autosize_excel(ARQUIVO_SAIDA_XLSX)

        print(f"\n✅ '{ARQUIVO_SAIDA_XLSX}' com {len(df)} registros.")
        print(f"   Datas preenchidas:  {df['data'].notna().sum()} / {len(df)}")
        print(f"   Lugar identificado: {(df['lugar'] != 'NÃO IDENTIFICADO').sum()} / {len(df)}")
        print(f"   OVOS capturados:    {(df['produto'] == 'OVOS').sum()}")

        print("\n📋 Prévia (primeiros 30 registros):")
        print()
        imprimir_tabela(df.head(30))
    else:
        print("\n⚠️ Nenhum dado.")

    unir_pdfs(caminhos, ARQUIVO_SAIDA_PDF)


if __name__ == "__main__":
    main()